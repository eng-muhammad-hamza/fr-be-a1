import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_report():
    doc = Document()

    # Style definitions
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Title
    title = doc.add_paragraph()
    title_run = title.add_run("W2 · A1 — Build Your First CRUD API")
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run("Comprehensive Engineering Report & AI Rematch Analysis\nBackend Development Program")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    doc.add_paragraph("─" * 60)

    # Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Architecture", level=1)
    doc.add_paragraph(
        "This project implements an industry-standard, production-aligned RESTful CRUD API using FastAPI and Python 3.14. "
        "The server models an in-memory task management engine (to-do list) supporting complete Create, Read, Update, and "
        "Delete operations, robust input validation, uniform error responses (400, 404), automatic interactive Swagger UI "
        "documentation at /docs, and extended production features including pagination, keyword search, filtering, and compute statistics."
    )

    # Endpoints Table
    h2 = doc.add_heading("2. API Specification & Endpoints", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Method"
    hdr_cells[1].text = "Endpoint"
    hdr_cells[2].text = "Description"
    hdr_cells[3].text = "Status Code"

    for cell in hdr_cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    endpoints_data = [
        ("GET", "/", "API metadata & version info", "200 OK"),
        ("GET", "/health", "Liveness & health monitoring probe", "200 OK"),
        ("GET", "/tasks", "List tasks (supports ?done, ?search, ?limit, ?offset)", "200 OK"),
        ("GET", "/tasks/{id}", "Retrieve single task by numeric ID", "200 OK / 404 Not Found"),
        ("POST", "/tasks", "Create new task with auto-increment ID", "201 Created / 400 Bad Request"),
        ("PUT", "/tasks/{id}", "Update task title and/or done status", "200 OK / 400 / 404"),
        ("DELETE", "/tasks/{id}", "Remove task by ID", "204 No Content / 404"),
        ("GET", "/stats", "Compute total, done, and open tasks counts", "200 OK"),
        ("POST", "/reset", "Reset in-memory storage to initial 3-task seed", "200 OK"),
    ]

    for method, path, desc, status_code in endpoints_data:
        row_cells = table.add_row().cells
        row_cells[0].text = method
        row_cells[1].text = path
        row_cells[2].text = desc
        row_cells[3].text = status_code

    doc.add_paragraph()

    # Stage 5 Documentation & Visual Swagger
    doc.add_heading("3. Swagger UI & Interactive Documentation", level=1)
    doc.add_paragraph(
        "Interactive Swagger UI is exposed at /docs, generated via OpenAPI 3.1. It enables developers and testers to "
        "visually test the entire CRUD cycle ('Try it out') without writing custom terminal scripts."
    )

    if os.path.exists("docs_screenshot.png"):
        doc.add_paragraph("Below is the captured interface of the running Swagger UI:")
        doc.add_picture("docs_screenshot.png", width=Inches(6.2))

    # The Mortality Experiment
    doc.add_heading("4. The Mortality Experiment (In-Memory vs Database)", level=1)
    doc.add_paragraph(
        "Observation:\n"
        "When new tasks are created during a server session, they exist only in RAM within the Python process heap. "
        "Upon restarting the Uvicorn server, all dynamically created and modified tasks are completely wiped out, "
        "and the state reverts to the initial hardcoded list.\n\n"
        "Technical Rationale:\n"
        "Volatile memory (RAM) is freed by the Operating System when a process terminates. Production backends require "
        "persistent storage (such as PostgreSQL, MySQL, SQLite, or Redis) where data is serialized and committed to disk "
        "or write-ahead logs, ensuring state survives crashes, deployments, and restarts."
    )

    # AI Rematch
    doc.add_heading("5. Stage 7 Bonus: The AI Rematch ('AI vs Me')", level=1)
    doc.add_paragraph(
        "In Stage 7, we wrote a specification prompt from memory and instructed an AI assistant to implement the API. "
        "The generated code was quarantined in 'ai-version/main.py' and tested against our verification suite."
    )

    doc.add_heading("Initial Prompt (From Memory):", level=2)
    doc.add_paragraph(
        "\"Write a FastAPI REST API in Python for managing a to-do list in memory. "
        "It should have endpoints to get root info, health status, list all tasks, get a task by id, "
        "create a task, update a task, and delete a task. Validate inputs and return proper HTTP status "
        "codes and JSON errors for not found and bad request.\""
    )

    doc.add_heading("Code Review & Concrete Differences:", level=2)
    
    diff_points = [
        ("1. HTTP 422 vs HTTP 400 on Validation:",
         "The AI relied on FastAPI's default Pydantic validation error handler which emits HTTP 422 Unprocessable Entity "
         "with schema validation trees. The assignment specification strictly requires HTTP 400 Bad Request with a clear JSON error. "
         "Our hand-crafted version registered a custom exception_handler(RequestValidationError) to standardize this."),
        ("2. Error JSON Payload Shape (detail vs error):",
         "The AI generated FastAPI's standard HTTPException which serializes errors as {'detail': '...'} instead of {'error': '...'}. "
         "Our hand-crafted version guarantees consistency for automated clients by formatting error responses as {'error': message}."),
        ("3. In-Memory ID Collision on Deletion:",
         "The AI assigned IDs using len(tasks) + 1. If tasks are deleted (e.g. deleting task 1 out of 3), the next task created would "
         "receive ID 3, colliding with existing task 3! Our hand-crafted version used max([t['id'] for t in tasks], default=0) + 1, "
         "preventing ID collisions."),
        ("4. Empty Request Body Validation on PUT:",
         "When an empty body {} was sent to PUT /tasks/1, the AI version succeeded silently with 200 without changing fields. "
         "Our hand-crafted version rejects empty update payloads with HTTP 400 ('Request body cannot be empty')."),
    ]

    for title, desc in diff_points:
        p = doc.add_paragraph()
        r1 = p.add_run(f"• {title} ")
        r1.bold = True
        p.add_run(desc)

    doc.add_heading("The Rematch:", level=2)
    doc.add_paragraph(
        "We improved our prompt by converting it into a rigorous engineering specification specifying error payload keys, "
        "exception handlers, and ID generation rules. The regenerated code in 'ai-version/rematch.py' passed 100% of the test suite."
    )

    # Verification & Test Coverage
    doc.add_heading("6. Automated Verification & Test Results", level=1)
    doc.add_paragraph(
        "A full automated regression suite was developed in 'tests/test_api.py' using Pytest and Starlette TestClient. "
        "13 test cases cover all CRUD paths, input validation edge cases, 404 scenarios, query parameter filters, pagination, "
        "and seed resets. All 13 tests pass with 100% pass rate."
    )

    doc.save("assignment_report.docx")
    print("Report assignment_report.docx generated successfully!")

if __name__ == "__main__":
    generate_report()
